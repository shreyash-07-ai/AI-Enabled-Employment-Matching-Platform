import os
import re
import sys
import csv
import json
import shutil
import hashlib
import logging
import time
import string
import traceback
from datetime import datetime
from pathlib import Path
from collections import defaultdict

import nltk
from nltk.corpus import stopwords
from nltk.tokenize import word_tokenize

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from pdfminer.high_level import extract_text as pdf_extract
import docx

# ── Optional: BM25 (pip install rank_bm25) ──────────────────────────
try:
    from rank_bm25 import BM25Okapi
    BM25_AVAILABLE = True
except ImportError:
    BM25_AVAILABLE = False

# ── NLTK bootstrap ──────────────────────────────────────────────────
logging.getLogger("pdfminer").setLevel(logging.ERROR)

for pkg in ["punkt", "stopwords", "punkt_tab"]:
    try:
        nltk.download(pkg, quiet=True)
    except Exception:
        pass

# ════════════════════════════════════════════════════════════════════
# CONFIGURATION  ← Edit this block or pass a config JSON file
# ════════════════════════════════════════════════════════════════════
DEFAULT_CONFIG = {
    "resume_folder":     r"C:\Users\Shreyash Musmade\Desktop\BE\Main\resumes",
    "requirement_file":  r"C:\Users\Shreyash Musmade\Desktop\BE\Main\requirements\JR2.txt",
    "output_folder":     r"C:\Users\Shreyash Musmade\Desktop\BE\Main",
    "top_n":             10,          # how many resumes to shortlist
    "min_score_pct":     13.0,      # ignore resumes below this % match
    "tfidf_weight":      0.6,         # weight for TF-IDF score (0-1)
    "bm25_weight":       0.4,         # weight for BM25 score  (ignored if BM25 unavailable)
    "copy_resumes":      True,        # copy shortlisted files to output folder
    "generate_html":     True,        # generate visual HTML report
    "generate_csv":      True,        # generate CSV report
    "log_level":         "INFO",      # DEBUG | INFO | WARNING | ERROR
}

# ════════════════════════════════════════════════════════════════════
# LOGGING SETUP
# ════════════════════════════════════════════════════════════════════
def setup_logging(level: str, log_file: str = None):
    handlers = [logging.StreamHandler(sys.stdout)]
    if log_file:
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))

    logging.basicConfig(
        level=getattr(logging, level.upper(), logging.INFO),
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%H:%M:%S",
        handlers=handlers,
        force=True,
    )

log = logging.getLogger(__name__)

# ════════════════════════════════════════════════════════════════════
# TEXT EXTRACTION
# ════════════════════════════════════════════════════════════════════
def extract_pdf(path: str) -> str:
    try:
        return pdf_extract(path) or ""
    except Exception as e:
        log.warning(f"PDF extraction failed [{path}]: {e}")
        return ""


def extract_docx(path: str) -> str:
    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs]
        # Also grab table text (often contains skills)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return " ".join(parts)
    except Exception as e:
        log.warning(f"DOCX extraction failed [{path}]: {e}")
        return ""


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    elif ext in (".docx", ".doc"):
        return extract_docx(path)
    else:
        log.warning(f"Unsupported file type: {path}")
        return ""

# ════════════════════════════════════════════════════════════════════
# TEXT CLEANING
# ════════════════════════════════════════════════════════════════════
STOP_WORDS = set(stopwords.words("english"))

def clean_text(text: str) -> str:
    text = text.lower()
    # Remove URLs and emails
    text = re.sub(r"http\S+|www\S+|\S+@\S+", " ", text)
    # Remove non-alphanumeric except spaces
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    tokens = word_tokenize(text)
    tokens = [w for w in tokens if w not in STOP_WORDS and len(w) > 1]
    return " ".join(tokens)


def tokenize(text: str) -> list:
    return text.split()

# ════════════════════════════════════════════════════════════════════
# KEYWORD ANALYSIS
# ════════════════════════════════════════════════════════════════════
def extract_keywords(text: str, top_n: int = 30) -> list:
    """Extract the most significant words from a text using TF-IDF on single doc."""
    vectorizer = TfidfVectorizer(max_features=top_n, ngram_range=(1, 2))
    try:
        mat = vectorizer.fit_transform([text])
        return vectorizer.get_feature_names_out().tolist()
    except Exception:
        return []


def keyword_overlap(resume_clean: str, jd_keywords: list) -> dict:
    """Returns matched and missing keywords."""
    resume_words = set(resume_clean.split())
    matched = [kw for kw in jd_keywords if all(w in resume_words for w in kw.split())]
    missing = [kw for kw in jd_keywords if kw not in matched]
    return {
        "matched": matched,
        "missing": missing,
        "match_pct": round(len(matched) / max(len(jd_keywords), 1) * 100, 1),
    }

# ════════════════════════════════════════════════════════════════════
# DUPLICATE DETECTION
# ════════════════════════════════════════════════════════════════════
def file_hash(path: str) -> str:
    h = hashlib.md5()
    try:
        with open(path, "rb") as f:
            h.update(f.read())
    except Exception:
        pass
    return h.hexdigest()

# ════════════════════════════════════════════════════════════════════
# SCORING
# ════════════════════════════════════════════════════════════════════
def tfidf_scores(resume_texts: list, jd_text: str) -> list:
    docs = resume_texts + [jd_text]
    vec = TfidfVectorizer()
    mat = vec.fit_transform(docs)
    jd_vec = mat[-1]
    res_vecs = mat[:-1]
    sims = cosine_similarity(jd_vec, res_vecs)[0]
    return sims.tolist()


def bm25_scores(resume_texts: list, jd_text: str) -> list:
    if not BM25_AVAILABLE:
        return [0.0] * len(resume_texts)
    corpus = [tokenize(t) for t in resume_texts]
    bm25 = BM25Okapi(corpus)
    query = tokenize(jd_text)
    raw = bm25.get_scores(query)
    max_score = max(raw) if max(raw) > 0 else 1
    return (raw / max_score).tolist()


def combined_scores(
    resume_texts: list,
    jd_text: str,
    tfidf_w: float = 0.6,
    bm25_w: float = 0.4,
) -> list:
    tf_scores = tfidf_scores(resume_texts, jd_text)
    bm_scores  = bm25_scores(resume_texts, jd_text)

    if not BM25_AVAILABLE:
        log.info("BM25 not available (pip install rank_bm25). Using TF-IDF only.")
        return tf_scores

    combined = [
        tfidf_w * tf + bm25_w * bm
        for tf, bm in zip(tf_scores, bm_scores)
    ]
    return combined

# ════════════════════════════════════════════════════════════════════
# FOLDER MANAGEMENT
# ════════════════════════════════════════════════════════════════════
def create_output_folder(base: str, req_file: str) -> str:
    name = Path(req_file).stem
    folder = Path(base) / name
    i = 1
    original = folder
    while folder.exists():
        folder = Path(f"{original}_{i}")
        i += 1
    folder.mkdir(parents=True)
    return str(folder)

# ════════════════════════════════════════════════════════════════════
# REPORT GENERATION
# ════════════════════════════════════════════════════════════════════
def write_csv(results: list, path: str):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "rank", "filename", "score_pct", "keyword_match_pct",
            "matched_keywords", "missing_keywords"
        ])
        writer.writeheader()
        for r in results:
            writer.writerow({
                "rank":               r["rank"],
                "filename":           r["name"],
                "score_pct":          r["score_pct"],
                "keyword_match_pct":  r["keyword"]["match_pct"],
                "matched_keywords":   ", ".join(r["keyword"]["matched"]),
                "missing_keywords":   ", ".join(r["keyword"]["missing"]),
            })
    log.info(f"CSV report saved: {path}")


def write_html(results: list, jd_path: str, path: str, elapsed: float):
    job_name = Path(jd_path).stem
    now      = datetime.now().strftime("%d %b %Y, %H:%M")

    def bar(pct, color="#4ade80"):
        return (
            f'<div style="background:#1e293b;border-radius:6px;height:10px;width:100%">'
            f'<div style="background:{color};width:{min(pct,100):.1f}%;height:10px;'
            f'border-radius:6px;transition:width .6s ease"></div></div>'
        )

    rows = ""
    for r in results:
        kw = r["keyword"]
        matched_tags = " ".join(
            f'<span style="background:#166534;color:#86efac;padding:2px 8px;'
            f'border-radius:99px;font-size:11px">{k}</span>'
            for k in kw["matched"][:12]
        )
        missing_tags = " ".join(
            f'<span style="background:#7f1d1d;color:#fca5a5;padding:2px 8px;'
            f'border-radius:99px;font-size:11px">{k}</span>'
            for k in kw["missing"][:12]
        )
        medal = {1: "🥇", 2: "🥈", 3: "🥉"}.get(r["rank"], f"#{r['rank']}")
        score_color = "#4ade80" if r["score_pct"] >= 60 else (
                      "#facc15" if r["score_pct"] >= 35 else "#f87171")

        rows += f"""
        <div style="background:#1e293b;border-radius:14px;padding:22px 26px;
                    margin-bottom:16px;border:1px solid #334155">
          <div style="display:flex;align-items:center;gap:14px;margin-bottom:12px">
            <span style="font-size:26px">{medal}</span>
            <div style="flex:1">
              <div style="font-size:16px;font-weight:700;color:#f1f5f9">{r['name']}</div>
              <div style="font-size:12px;color:#64748b;margin-top:2px">
                Combined score: <b style="color:{score_color}">{r['score_pct']}%</b>
                &nbsp;|&nbsp; Keyword match: <b style="color:#a78bfa">{kw['match_pct']}%</b>
              </div>
            </div>
            <div style="font-size:28px;font-weight:900;color:{score_color}">{r['score_pct']}%</div>
          </div>
          {bar(r['score_pct'], score_color)}
          <div style="margin-top:14px">
            <div style="font-size:11px;color:#64748b;margin-bottom:6px">✅ MATCHED SKILLS</div>
            {matched_tags or '<span style="color:#475569;font-size:12px">None detected</span>'}
          </div>
          <div style="margin-top:10px">
            <div style="font-size:11px;color:#64748b;margin-bottom:6px">❌ MISSING SKILLS</div>
            {missing_tags or '<span style="color:#475569;font-size:12px">All key skills present</span>'}
          </div>
        </div>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Resume Screener Report — {job_name}</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: 'Segoe UI', system-ui, sans-serif;
          background: #0f172a; color: #cbd5e1; padding: 40px 20px; }}
  .container {{ max-width: 860px; margin: 0 auto; }}
  h1 {{ font-size: 28px; color: #f1f5f9; margin-bottom: 4px; }}
  .meta {{ font-size: 13px; color: #475569; margin-bottom: 32px; }}
  .stat {{ display:inline-block; background:#1e293b; border-radius:10px;
           padding:12px 20px; margin-right:12px; margin-bottom:20px; }}
  .stat-val {{ font-size:22px; font-weight:800; color:#a78bfa; }}
  .stat-lbl {{ font-size:11px; color:#64748b; margin-top:2px; }}
</style>
</head>
<body>
<div class="container">
  <h1>🎯 Resume Screener Report</h1>
  <div class="meta">Job: <b style="color:#f1f5f9">{job_name}</b> &nbsp;|&nbsp;
       Generated: {now} &nbsp;|&nbsp; Time: {elapsed:.2f}s</div>

  <div>
    <div class="stat">
      <div class="stat-val">{len(results)}</div>
      <div class="stat-lbl">Shortlisted</div>
    </div>
    <div class="stat">
      <div class="stat-val">{results[0]['score_pct'] if results else 0}%</div>
      <div class="stat-lbl">Top Score</div>
    </div>
    <div class="stat">
      <div class="stat-val">{'BM25 + TF-IDF' if BM25_AVAILABLE else 'TF-IDF'}</div>
      <div class="stat-lbl">Scoring Model</div>
    </div>
  </div>

  <div style="margin-top:10px">{rows}</div>
  <div style="font-size:11px;color:#334155;margin-top:30px;text-align:center">
    Generated by AI Resume Screener v2.0
  </div>
</div>
</body>
</html>"""

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    log.info(f"HTML report saved: {path}")

# ════════════════════════════════════════════════════════════════════
# MAIN
# ════════════════════════════════════════════════════════════════════
def load_config(config_path: str = None) -> dict:
    cfg = DEFAULT_CONFIG.copy()
    if config_path and Path(config_path).exists():
        with open(config_path, "r") as f:
            cfg.update(json.load(f))
    return cfg


def main(config_path: str = None):
    cfg = load_config(config_path)

    # ── Setup logging ────────────────────────────────────────────────
    log_file = os.path.join(cfg["output_folder"], "screener.log")
    setup_logging(cfg["log_level"], log_file)

    start = time.time()

    log.info("=" * 60)
    log.info("  AI RESUME SCREENER v2.0  STARTED")
    log.info("=" * 60)
    log.info(f"Resume folder   : {cfg['resume_folder']}")
    log.info(f"Requirement file: {cfg['requirement_file']}")
    log.info(f"Top N           : {cfg['top_n']}")
    log.info(f"Min score       : {cfg['min_score_pct']}%")
    log.info(f"BM25 available  : {BM25_AVAILABLE}")

    # ── Load JD ─────────────────────────────────────────────────────
    if not Path(cfg["requirement_file"]).exists():
        log.error(f"Requirement file not found: {cfg['requirement_file']}")
        sys.exit(1)

    with open(cfg["requirement_file"], "r", encoding="utf-8") as f:
        raw_jd = f.read()
    jd_clean   = clean_text(raw_jd)
    jd_keywords = extract_keywords(jd_clean, top_n=30)
    log.info(f"JD keywords extracted: {len(jd_keywords)}")

    # ── Load resumes ─────────────────────────────────────────────────
    resume_dir = Path(cfg["resume_folder"])
    if not resume_dir.exists():
        log.error(f"Resume folder not found: {resume_dir}")
        sys.exit(1)

    all_files = [
        str(p) for p in resume_dir.iterdir()
        if p.suffix.lower() in (".pdf", ".docx", ".doc")
    ]
    log.info(f"Found {len(all_files)} resume file(s)")

    # ── Extract & deduplicate ────────────────────────────────────────
    seen_hashes = {}
    resume_files, resume_texts = [], []
    skipped_dup = 0

    for path in all_files:
        h = file_hash(path)
        if h in seen_hashes:
            log.warning(f"Duplicate skipped: {Path(path).name} (same as {seen_hashes[h]})")
            skipped_dup += 1
            continue
        seen_hashes[h] = Path(path).name

        text = extract_text(path)
        if not text.strip():
            log.warning(f"No text extracted, skipping: {Path(path).name}")
            continue

        resume_files.append(path)
        resume_texts.append(clean_text(text))

    log.info(f"Processable resumes: {len(resume_files)} | Duplicates skipped: {skipped_dup}")

    if not resume_files:
        log.error("No valid resumes found. Exiting.")
        sys.exit(1)

    # ── Score ────────────────────────────────────────────────────────
    log.info("Scoring resumes...")
    scores = combined_scores(
        resume_texts, jd_clean,
        tfidf_w=cfg["tfidf_weight"],
        bm25_w=cfg["bm25_weight"],
    )

    # ── Build results ────────────────────────────────────────────────
    results = []
    for i, (path, score) in enumerate(zip(resume_files, scores)):
        score_pct = round(score * 100, 2)
        if score_pct < cfg["min_score_pct"]:
            continue
        kw = keyword_overlap(resume_texts[i], jd_keywords)
        results.append({
            "file":      path,
            "name":      Path(path).name,
            "score":     score,
            "score_pct": score_pct,
            "keyword":   kw,
        })

    ranked = sorted(results, key=lambda x: x["score"], reverse=True)
    top    = ranked[: cfg["top_n"]]

    for i, item in enumerate(top, start=1):
        item["rank"] = i

    # ── Print summary ─────────────────────────────────────────────────
    print("\n" + "═" * 60)
    print(f"  TOP {len(top)} SHORTLISTED RESUMES")
    print("═" * 60)
    for item in top:
        kw = item["keyword"]
        print(
            f"  Rank {item['rank']:>2} │ {item['score_pct']:>6.2f}% │ "
            f"KW {kw['match_pct']:>5.1f}% │ {item['name']}"
        )
    print("═" * 60)
    print(f"  Total scored: {len(results)} | Below threshold: "
          f"{len(resume_files) - len(results)}")
    print()

    # ── Save files ────────────────────────────────────────────────────
    save_folder = create_output_folder(cfg["output_folder"], cfg["requirement_file"])
    log.info(f"Output folder: {save_folder}")

    if cfg["copy_resumes"]:
        for item in top:
            dest = os.path.join(save_folder, f"Rank_{item['rank']}_{item['name']}")
            shutil.copy(item["file"], dest)
        log.info(f"Copied {len(top)} resumes to output folder")

    if cfg["generate_csv"]:
        csv_path = os.path.join(save_folder, "report.csv")
        write_csv(top, csv_path)

    elapsed = round(time.time() - start, 2)

    if cfg["generate_html"]:
        html_path = os.path.join(save_folder, "report.html")
        write_html(top, cfg["requirement_file"], html_path, elapsed)

    log.info(f"Done in {elapsed}s")
    print(f"✅ Results saved to: {save_folder}\n")


if __name__ == "__main__":
    # Optional: pass path to a JSON config file as argument
    config_file = sys.argv[1] if len(sys.argv) > 1 else None
    main(config_file)